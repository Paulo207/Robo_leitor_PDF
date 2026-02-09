"""
Módulo para escrita de dados em planilhas Excel
"""

import pandas as pd
from pathlib import Path
import logging
from typing import Dict, List
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger(__name__)


class WriterExcel:
    """Classe para escrever dados extraídos em Excel"""
    
    def __init__(self, arquivo_excel: str = "faturas_coelba.xlsx"):
        self.arquivo_excel = arquivo_excel
        self.colunas = [
            'Cliente', 'UC', 'Mês Referência', 'Data Vencimento',
            'Valor Total', 'Consumo kWh', 'ICMS', 'PIS', 'COFINS', 'CIP',
            'Data Processamento', 'Arquivo PDF'
        ]
    
    def criar_planilha_modelo(self):
        """Cria arquivo Excel modelo com cabeçalhos formatados"""
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Faturas"
            
            # Adicionar cabeçalhos
            for col, nome in enumerate(self.colunas, start=1):
                celula = ws.cell(row=1, column=col, value=nome)
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                celula.alignment = Alignment(horizontal="center", vertical="center")
            
            # Ajustar largura das colunas
            larguras = [30, 15, 15, 15, 15, 12, 12, 12, 12, 12, 18, 30]
            for col, largura in enumerate(larguras, start=1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = largura
            
            # Salvar
            wb.save(self.arquivo_excel)
            logger.info(f"Planilha modelo criada: {self.arquivo_excel}")
            
        except Exception as e:
            logger.error(f"Erro ao criar planilha modelo: {str(e)}")
    
    def adicionar_fatura(self, dados: Dict, arquivo_pdf: str = "") -> bool:
        """
        Adiciona uma nova linha com dados da fatura
        
        Args:
            dados: Dicionário com campos extraídos
            arquivo_pdf: Nome do arquivo PDF processado
            
        Returns:
            True se sucesso, False caso contrário
        """
        try:
            # Criar arquivo se não existir
            if not Path(self.arquivo_excel).exists():
                self.criar_planilha_modelo()
            
            # Preparar linha de dados
            linha = {
                'Cliente': dados.get('cliente', ''),
                'UC': dados.get('uc', ''),
                'Mês Referência': dados.get('mes_referencia', ''),
                'Data Vencimento': dados.get('data_vencimento', ''),
                'Valor Total': dados.get('valor_total', ''),
                'Consumo kWh': dados.get('consumo_kwh', ''),
                'ICMS': dados.get('impostos', {}).get('icms', ''),
                'PIS': dados.get('impostos', {}).get('pis', ''),
                'COFINS': dados.get('impostos', {}).get('cofins', ''),
                'CIP': dados.get('impostos', {}).get('cip', ''),
                'Data Processamento': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
                'Arquivo PDF': arquivo_pdf
            }
            
            # Ler arquivo existente
            df = pd.read_excel(self.arquivo_excel)
            
            # Adicionar nova linha
            df = pd.concat([df, pd.DataFrame([linha])], ignore_index=True)
            
            # Salvar
            df.to_excel(self.arquivo_excel, index=False)
            logger.info(f"Fatura adicionada ao Excel: UC {dados.get('uc', 'N/A')}")
            
            return True
            
        except Exception as e:
            logger.error(f"Erro ao adicionar fatura ao Excel: {str(e)}")
            return False
    
    def adicionar_multiplas_faturas(self, lista_dados: List[Dict]) -> int:
        """
        Adiciona múltiplas faturas de uma vez
        
        Args:
            lista_dados: Lista de dicionários com dados das faturas
            
        Returns:
            Número de faturas adicionadas com sucesso
        """
        contador = 0
        for dados in lista_dados:
            if self.adicionar_fatura(dados):
                contador += 1
        
        logger.info(f"{contador}/{len(lista_dados)} faturas adicionadas ao Excel")
        return contador
    
    def verificar_duplicata(self, uc: str, mes_referencia: str) -> bool:
        """
        Verifica se já existe fatura para esta UC e mês
        
        Args:
            uc: Número da unidade consumidora
            mes_referencia: Mês de referência
            
        Returns:
            True se já existe, False caso contrário
        """
        try:
            if not Path(self.arquivo_excel).exists():
                return False
            
            df = pd.read_excel(self.arquivo_excel)
            
            # Buscar por UC e mês
            duplicata = df[
                (df['UC'].astype(str) == str(uc)) &
                (df['Mês Referência'].astype(str) == str(mes_referencia))
            ]
            
            if len(duplicata) > 0:
                logger.warning(f"Duplicata encontrada: UC {uc}, Mês {mes_referencia}")
                return True
            
            return False
            
        except Exception as e:
            logger.error(f"Erro ao verificar duplicata: {str(e)}")
            return False
    
    def obter_estatisticas(self) -> Dict:
        """
        Obtém estatísticas das faturas processadas
        
        Returns:
            Dicionário com estatísticas
        """
        try:
            if not Path(self.arquivo_excel).exists():
                return {}
            
            df = pd.read_excel(self.arquivo_excel)
            
            # Converter valores para numérico (remover R$, vírgulas, etc)
            df['Valor Total Num'] = df['Valor Total'].apply(self._converter_para_numero)
            df['Consumo Num'] = df['Consumo kWh'].apply(self._converter_para_numero)
            
            stats = {
                'total_faturas': len(df),
                'valor_total_soma': df['Valor Total Num'].sum(),
                'consumo_total_soma': df['Consumo Num'].sum(),
                'valor_medio': df['Valor Total Num'].mean(),
                'consumo_medio': df['Consumo Num'].mean(),
                'clientes_unicos': df['Cliente'].nunique(),
                'ucs_unicas': df['UC'].nunique()
            }
            
            return stats
            
        except Exception as e:
            logger.error(f"Erro ao obter estatísticas: {str(e)}")
            return {}
    
    def _converter_para_numero(self, valor: str) -> float:
        """Converte string com valor monetário para float"""
        try:
            if pd.isna(valor) or valor == '':
                return 0.0
            
            # Remover R$, espaços
            valor_limpo = str(valor).replace('R$', '').replace(' ', '')
            # Remover pontos (milhares) e substituir vírgula por ponto
            valor_limpo = valor_limpo.replace('.', '').replace(',', '.')
            
            return float(valor_limpo)
            
        except:
            return 0.0
